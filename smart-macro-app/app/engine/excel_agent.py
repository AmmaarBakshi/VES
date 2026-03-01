import pandas as pd
import os
from app.ai.llm_client import query_ollama
from app.ai.prompts import get_prompt
from langchain_ollama import OllamaLLM
from app.engine.cot_engine import stream_cot_plan, make_excel_plan_prompt, make_enhance_plan_prompt
from app.utils.config_loader import load_ai_config

# Shared LLM for CoT planning
_plan_llm = None
def _get_plan_llm():
    global _plan_llm
    if _plan_llm is None:
        config = load_ai_config()
        model_name = config.get("active_model", "llama3.2")
        _plan_llm = OllamaLLM(model=model_name, temperature=0.5)
    return _plan_llm


def process_excel_file(file_path, instruction, thought_callback=None):
    """
    Reads an Excel file, applies AI to the first column, and saves the result.
    thought_callback: optional callable(str) for streaming CoT reasoning.
    """
    try:
        print(f"Reading Excel: {file_path}")
        df = pd.read_excel(file_path)
        col_name = df.columns[0]

        # ── Agentic CoT: Stream planning thoughts before processing ──
        if thought_callback:
            thought_callback(f"◆ File: {os.path.basename(file_path)}\n\n")
            col_names = list(df.columns)
            plan_prompt = make_excel_plan_prompt(instruction, len(df), col_names)
            stream_cot_plan(_get_plan_llm(), plan_prompt, thought_callback)
            thought_callback("\n\n─── Processing Cells ───\n\n")
        
        results = []
        for val in df[col_name]:
            if pd.isna(val) or str(val).strip() == "":
                results.append("")
                continue
            prompt = get_prompt(instruction, str(val))
            resp = query_ollama(prompt)
            results.append(resp)

        df['AI_Output'] = results
        
        base_name = os.path.basename(file_path)
        dir_name = os.path.dirname(file_path)
        new_path = os.path.join(dir_name, f"PROCESSED_{base_name}")
        
        df.to_excel(new_path, index=False)
        return new_path
    except Exception as e:
        print(f"Excel Error: {e}")
        return f"Error: {e}"


def enrich_excel_file(file_path, options, style="professional", thought_callback=None):
    """
    Enhance an existing Excel file with AI-powered improvements.
    thought_callback: optional callable(str) for streaming CoT reasoning.
    """
    try:
        from app.engine.content_enricher import ContentEnricher
        from openpyxl import load_workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        import time
        
        enricher = ContentEnricher()
        
        print(f"Reading Excel: {file_path}")
        df = pd.read_excel(file_path)
        text_columns = df.select_dtypes(include=['object']).columns.tolist()

        # ── Agentic CoT: Stream planning thoughts before enhancement ──
        if thought_callback:
            thought_callback(f"◆ Enhancing: {os.path.basename(file_path)}\n\n")
            plan_prompt = make_enhance_plan_prompt("Excel spreadsheet (.xlsx)", options, style)
            stream_cot_plan(_get_plan_llm(), plan_prompt, thought_callback)
            thought_callback("\n\n─── Applying Enhancements ───\n\n")
        
        # Improve content in text columns
        if options.get('improve_content', False) and text_columns:
            print("Improving cell content...")
            for col in text_columns:
                improved_values = []
                for val in df[col]:
                    if pd.isna(val) or str(val).strip() == "":
                        improved_values.append(val)
                    else:
                        improved = enricher.improve_text(str(val), style=style, intensity="light")
                        improved_values.append(improved)
                df[col] = improved_values
        
        # Fix consistency
        if options.get('fix_consistency', False) and text_columns:
            print("Ensuring consistency...")
            for col in text_columns:
                values = [str(v) if not pd.isna(v) else "" for v in df[col]]
                if values:
                    consistent_values = enricher.check_consistency(values, target_style=style)
                    df[col] = consistent_values
        
        # Add summary row
        if options.get('add_summaries', False):
            print("Adding summary insights...")
            summary_row = {}
            for col in df.columns:
                if col in text_columns:
                    all_text = " ".join([str(v) for v in df[col] if not pd.isna(v)])
                    if all_text.strip():
                        summary = enricher.summarize_content(all_text[:500], max_length=30)
                        summary_row[col] = f"Summary: {summary}" if summary else ""
                    else:
                        summary_row[col] = ""
                else:
                    try:
                        summary_row[col] = f"Total: {df[col].sum():.2f}"
                    except:
                        summary_row[col] = ""
            df = pd.concat([df, pd.DataFrame([summary_row])], ignore_index=True)
        
        base_name = os.path.basename(file_path)
        name_without_ext = os.path.splitext(base_name)[0]
        dir_name = os.path.dirname(file_path)
        temp_path = os.path.join(dir_name, f"temp_{int(time.time())}.xlsx")
        df.to_excel(temp_path, index=False)
        
        if options.get('auto_format', False):
            print("Applying formatting...")
            wb = load_workbook(temp_path)
            ws = wb.active
            
            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF", size=12)
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            
            thin_border = Border(
                left=Side(style='thin'), right=Side(style='thin'),
                top=Side(style='thin'), bottom=Side(style='thin')
            )
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
                for cell in row:
                    cell.border = thin_border
                    if cell.row > 1:
                        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width
            
            final_path = os.path.join(dir_name, f"{name_without_ext}_ENHANCED_{int(time.time())}.xlsx")
            wb.save(final_path)
            try:
                os.remove(temp_path)
            except:
                pass
        else:
            final_path = os.path.join(dir_name, f"{name_without_ext}_ENHANCED_{int(time.time())}.xlsx")
            os.rename(temp_path, final_path)
        
        print(f"✅ Enhanced Excel saved: {final_path}")
        return final_path
        
    except Exception as e:
        print(f"Excel Enrichment Error: {e}")
        return f"Error: {e}"


def suggest_excel_improvements(file_path):
    """Analyze Excel file and provide improvement suggestions."""
    try:
        from app.engine.content_enricher import ContentEnricher
        
        enricher = ContentEnricher()
        df = pd.read_excel(file_path)
        
        suggestions = []
        suggestions.append("=== Excel Data Analysis ===\n")
        suggestions.append(f"Total Rows: {len(df)}")
        suggestions.append(f"Total Columns: {len(df.columns)}")
        suggestions.append("")
        
        empty_cells = df.isnull().sum().sum()
        if empty_cells > 0:
            suggestions.append(f"⚠️ Found {empty_cells} empty cells - consider filling or removing")
        
        text_columns = df.select_dtypes(include=['object']).columns.tolist()
        if text_columns:
            suggestions.append(f"\n📝 Text columns found: {', '.join(text_columns)}")
            if len(df) > 0:
                sample_text = " ".join([str(v) for v in df[text_columns[0]].head(3) if not pd.isna(v)])
                if sample_text:
                    ai_suggestions = enricher.suggest_improvements(
                        sample_text,
                        context="Excel spreadsheet data"
                    )
                    if ai_suggestions:
                        suggestions.append("\n💡 Content Improvement Suggestions:")
                        suggestions.extend(ai_suggestions)
        
        return suggestions
        
    except Exception as e:
        return [f"Error analyzing Excel: {e}"]
