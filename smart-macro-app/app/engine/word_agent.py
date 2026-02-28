import os
from docx import Document
from app.ai.llm_client import query_ollama
from langchain_ollama import OllamaLLM
from app.engine.cot_engine import stream_cot_plan, make_word_plan_prompt, make_enhance_plan_prompt

# Shared LLM for CoT planning (only used for streaming the plan, not document edits)
_plan_llm = None
def _get_plan_llm():
    global _plan_llm
    if _plan_llm is None:
        _plan_llm = OllamaLLM(model="llama3", temperature=0.5)
    return _plan_llm


def process_word_document(file_path, instruction="Fix grammar and make professional", thought_callback=None):
    """
    Reads a .docx file, processes each paragraph with AI, and saves a new copy.
    thought_callback: optional callable(str) for streaming CoT reasoning.
    """
    try:
        doc = Document(file_path)
        
        total_paragraphs = len([p for p in doc.paragraphs if p.text.strip()])
        processed_count = 0

        # ── Agentic CoT: Stream planning thoughts before processing ──
        if thought_callback:
            thought_callback(f"◆ File: {os.path.basename(file_path)}\n\n")
            plan_prompt = make_word_plan_prompt(instruction, total_paragraphs)
            stream_cot_plan(_get_plan_llm(), plan_prompt, thought_callback)
            thought_callback("\n\n─── Processing Paragraphs ───\n\n")

        print(f"Processing {total_paragraphs} paragraphs...")

        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                full_prompt = f"Instruction: {instruction}\n\nText: {paragraph.text}"
                new_text = query_ollama(full_prompt)
                paragraph.text = new_text
                processed_count += 1
                print(f" - Processed paragraph {processed_count}/{total_paragraphs}")

        dir_name = os.path.dirname(file_path)
        base_name = os.path.basename(file_path)
        new_path = os.path.join(dir_name, f"PROCESSED_{base_name}")
        
        doc.save(new_path)
        return new_path

    except Exception as e:
        print(f"Error processing file: {e}")
        return None


def enrich_word_document(file_path, options, style="professional", thought_callback=None):
    """
    Enhance an existing Word document with AI-powered improvements.
    thought_callback: optional callable(str) for streaming CoT reasoning.
    """
    try:
        from app.engine.content_enricher import ContentEnricher
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        import time
        
        enricher = ContentEnricher()
        doc = Document(file_path)
        all_paragraphs = [p for p in doc.paragraphs if p.text.strip()]

        # ── Agentic CoT: Stream planning thoughts before enhancement ──
        if thought_callback:
            file_type = "Word document (.docx)"
            thought_callback(f"◆ Enhancing: {os.path.basename(file_path)}\n\n")
            plan_prompt = make_enhance_plan_prompt(file_type, options, style)
            stream_cot_plan(_get_plan_llm(), plan_prompt, thought_callback)
            thought_callback("\n\n─── Applying Enhancements ───\n\n")

        # Improve paragraphs (batched for speed)
        if options.get('improve_paragraphs', False):
            print(f"Improving {len(all_paragraphs)} paragraphs (batched)...")
            BATCH_SIZE = 5
            for batch_start in range(0, len(all_paragraphs), BATCH_SIZE):
                batch = all_paragraphs[batch_start:batch_start + BATCH_SIZE]
                batch_texts = [p.text for p in batch if p.text.strip()]
                if batch_texts:
                    improved = enricher.improve_text_batch(batch_texts, style=style, intensity="moderate")
                    text_idx = 0
                    for p in batch:
                        if p.text.strip() and text_idx < len(improved):
                            p.text = improved[text_idx]
                            text_idx += 1
                    print(f" - Batch {batch_start//BATCH_SIZE + 1}: processed {len(batch_texts)} paragraphs")
        
        # Fix consistency
        if options.get('fix_consistency', False):
            print("Ensuring style consistency...")
            paragraph_texts = [p.text for p in all_paragraphs if p.text.strip()]
            if paragraph_texts:
                consistent_texts = enricher.check_consistency(paragraph_texts, target_style=style)
                text_idx = 0
                for paragraph in doc.paragraphs:
                    if paragraph.text.strip() and text_idx < len(consistent_texts):
                        paragraph.text = consistent_texts[text_idx]
                        text_idx += 1
        
        # Add document summary
        if options.get('add_summary', False):
            print("Generating document summary...")
            full_text = " ".join([p.text for p in all_paragraphs])
            if full_text.strip():
                summary = enricher.summarize_content(full_text[:2000], max_length=150)
                if summary:
                    summary_para = doc.paragraphs[0].insert_paragraph_before("Executive Summary")
                    summary_para.style = 'Heading 1'
                    content_para = doc.paragraphs[1].insert_paragraph_before(summary)
                    content_para.style = 'Normal'
                    doc.paragraphs[2].insert_paragraph_before("")
        
        # Apply formatting
        if options.get('auto_format', False):
            print("Applying formatting...")
            for paragraph in doc.paragraphs:
                for run in paragraph.runs:
                    run.font.name = 'Calibri'
                    if run.font.size is None or run.font.size < Pt(11):
                        run.font.size = Pt(11)
                paragraph.paragraph_format.line_spacing = 1.15
                paragraph.paragraph_format.space_after = Pt(6)
        
        base_name = os.path.basename(file_path)
        name_without_ext = os.path.splitext(base_name)[0]
        dir_name = os.path.dirname(file_path)
        import time as _time
        new_path = os.path.join(dir_name, f"{name_without_ext}_ENHANCED_{int(_time.time())}.docx")
        
        doc.save(new_path)
        print(f"✅ Enhanced document saved: {new_path}")
        return new_path
        
    except Exception as e:
        print(f"Word Enrichment Error: {e}")
        return None


def suggest_word_improvements(file_path):
    """Analyze Word document and provide improvement suggestions."""
    try:
        from app.engine.content_enricher import ContentEnricher
        
        enricher = ContentEnricher()
        doc = Document(file_path)
        
        suggestions = []
        suggestions.append("=== Document Analysis ===\n")
        
        paragraphs = [p for p in doc.paragraphs if p.text.strip()]
        suggestions.append(f"Total Paragraphs: {len(paragraphs)}")
        
        total_words = sum(len(p.text.split()) for p in paragraphs)
        suggestions.append(f"Total Words: {total_words}")
        suggestions.append("")
        
        if paragraphs:
            sample_text = " ".join([p.text for p in paragraphs[:3]])
            if sample_text:
                ai_suggestions = enricher.suggest_improvements(
                    sample_text,
                    context="Word document"
                )
                if ai_suggestions:
                    suggestions.append("💡 Content Improvement Suggestions:")
                    suggestions.extend(ai_suggestions)
        
        return suggestions
        
    except Exception as e:
        return [f"Error analyzing document: {e}"]
